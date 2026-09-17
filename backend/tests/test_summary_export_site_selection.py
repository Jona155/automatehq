import unittest
import uuid
from io import BytesIO

from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

from backend.app import create_app, db
from backend.app.auth_utils import encode_auth_token
from backend.app.models.business import Business
from backend.app.models.sites import Site
from backend.app.models.users import User
from backend.app.repositories.business_repository import BusinessRepository


class SummaryExportSiteSelectionTests(unittest.TestCase):
    def setUp(self):
        self.app = create_app()
        self.client = self.app.test_client()
        self.ctx = self.app.app_context()
        self.ctx.push()

        suffix = uuid.uuid4().hex[:8]
        self.business = Business(
            name=f'Summary Export {suffix}',
            code=f'summary-export-{suffix}',
            is_active=True,
        )
        self.other_business = Business(
            name=f'Other Summary Export {suffix}',
            code=f'other-summary-export-{suffix}',
            is_active=True,
        )
        db.session.add_all([self.business, self.other_business])
        db.session.flush()

        self.admin = User(
            business_id=self.business.id,
            full_name='Summary Export Admin',
            email=f'summary-export-{suffix}@example.com',
            role='ADMIN',
            password_hash='n/a',
            is_active=True,
        )
        self.site_a = Site(
            business_id=self.business.id,
            site_name=f'Alpha {suffix}',
            is_active=True,
        )
        self.site_b = Site(
            business_id=self.business.id,
            site_name=f'Beta {suffix}',
            is_active=True,
        )
        self.inactive_site = Site(
            business_id=self.business.id,
            site_name=f'Inactive {suffix}',
            is_active=False,
        )
        self.foreign_site = Site(
            business_id=self.other_business.id,
            site_name=f'Foreign {suffix}',
            is_active=True,
        )
        db.session.add_all([
            self.admin,
            self.site_a,
            self.site_b,
            self.inactive_site,
            self.foreign_site,
        ])
        db.session.commit()

        self.business_id = self.business.id
        self.other_business_id = self.other_business.id
        self.headers = {'Authorization': f'Bearer {encode_auth_token(self.admin.id)}'}
        self.base_url = '/api/sites/summary/export-batch?processing_month=2026-02-01'

    def tearDown(self):
        try:
            BusinessRepository().hard_delete(self.business_id)
            BusinessRepository().hard_delete(self.other_business_id)
        except Exception:
            db.session.rollback()
        self.ctx.pop()

    def _sheet_names(self, response):
        return load_workbook(BytesIO(response.data), read_only=True).sheetnames

    def test_omitted_site_ids_exports_all_active_sites(self):
        response = self.client.get(self.base_url, headers=self.headers)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._sheet_names(response), [self.site_a.site_name, self.site_b.site_name])
        self.assertIn('monthly_summary_all_sites_2026-02.xlsx', response.headers['Content-Disposition'])

    def test_selected_site_ids_export_only_selected_sheets(self):
        response = self.client.get(
            f'{self.base_url}&site_ids={self.site_b.id},{self.site_b.id}',
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._sheet_names(response), [self.site_b.site_name])
        self.assertIn('monthly_summary_selected_sites_2026-02.xlsx', response.headers['Content-Disposition'])

    def test_selected_inactive_site_respects_default_active_filter(self):
        response = self.client.get(
            f'{self.base_url}&site_ids={self.inactive_site.id}',
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._sheet_names(response), ['Sites'])

    def test_invalid_site_id_is_rejected(self):
        response = self.client.get(f'{self.base_url}&site_ids=not-a-uuid', headers=self.headers)

        self.assertEqual(response.status_code, 400)

    def test_site_from_another_business_is_rejected(self):
        response = self.client.get(
            f'{self.base_url}&site_ids={self.foreign_site.id}',
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 400)


if __name__ == '__main__':
    unittest.main()
