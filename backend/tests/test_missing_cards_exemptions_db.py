"""End-to-end tests for the missing-cards exemption rules against the DB.

Exercises ``compute_missing`` (and the report/group helpers it feeds) with real
work-card rows, so the extended cards subquery — single-card upload timestamp
and the manual full-approval flag — is verified, not just the pure classifier.

Requires a reachable local Postgres (same as the other DB-backed tests).
"""
import unittest
import uuid
from datetime import date, datetime, timezone

from dotenv import load_dotenv

load_dotenv()

from backend.app import create_app, db
from backend.app.models.business import Business
from backend.app.models.sites import Site, Employee
from backend.app.models.users import User
from backend.app.models.work_cards import WorkCard
from backend.app.services import missing_cards_service as mcs

MONTH = date(2026, 6, 1)          # June 2026, last day = 30
AFTER_MONTH_END = date(2026, 7, 2)  # "today" once the month has closed


def _utc(y, m, d, hh=12, mm=0):
    return datetime(y, m, d, hh, mm, tzinfo=timezone.utc)


class MissingCardsExemptionDBTests(unittest.TestCase):
    def setUp(self):
        self.app = create_app()
        self.ctx = self.app.app_context()
        self.ctx.push()

        self.business = Business(
            name='Exemption Test Biz',
            code=f'exempt-{uuid.uuid4().hex[:12]}',
            expected_work_cards_per_month=2,
        )
        db.session.add(self.business)
        db.session.flush()

        self.site = Site(business_id=self.business.id, site_name=f'Site {uuid.uuid4().hex[:8]}')
        db.session.add(self.site)
        db.session.flush()
        db.session.commit()

    def tearDown(self):
        try:
            WorkCard.query.filter_by(business_id=self.business.id).delete()
            Employee.query.filter_by(business_id=self.business.id).delete()
            Site.query.filter_by(business_id=self.business.id).update({'field_manager_id': None})
            User.query.filter_by(business_id=self.business.id).delete()
            Site.query.filter_by(business_id=self.business.id).delete()
            db.session.delete(db.session.get(Business, self.business.id))
            db.session.commit()
        except Exception:
            db.session.rollback()
        self.ctx.pop()

    # ---- helpers ----
    def _employee(self):
        emp = Employee(
            business_id=self.business.id,
            site_id=self.site.id,
            full_name=f'Emp {uuid.uuid4().hex[:8]}',
            passport_id=f'P{uuid.uuid4().hex[:10]}',
            is_active=True,
        )
        db.session.add(emp)
        db.session.flush()
        return emp

    def _card(self, emp, created_at, *, source='ADMIN_SINGLE',
              review_status='NEEDS_REVIEW', approved_through_day=None,
              completes_month=False, processing_month=MONTH):
        card = WorkCard(
            business_id=self.business.id,
            site_id=self.site.id,
            employee_id=emp.id,
            processing_month=processing_month,
            source=source,
            review_status=review_status,
            approved_through_day=approved_through_day,
            completes_month=completes_month,
            created_at=created_at,
        )
        db.session.add(card)
        db.session.flush()
        return card

    def _status(self, emp):
        rows = mcs.compute_missing(self.business.id, MONTH, today=AFTER_MONTH_END)
        row = next(r for r in rows if r['employee_id'] == str(emp.id))
        return row['status']

    # ---- Rule 1: late single card ----
    def test_single_card_on_last_day_is_complete(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 30))
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_COMPLETE)

    def test_single_card_on_fifth_is_complete(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 7, 5))
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_COMPLETE)

    def test_single_card_on_sixth_is_partial(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 7, 6))
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_PARTIAL)

    def test_single_mid_month_card_is_partial(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15))
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_PARTIAL)

    def test_two_cards_one_late_uses_normal_counting(self):
        # Rule 1 requires exactly one card; a mid-month + a late card is 2 -> COMPLETE
        # via the ordinary threshold, not the exemption. A mid-month + one-more
        # still-short case is covered by the single-card tests, so here two cards
        # simply satisfy expected=2.
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15))
        self._card(emp, _utc(2026, 6, 30))
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_COMPLETE)

    # ---- Rule 2: manual full-month approval ----
    def test_manual_approved_through_month_end_is_complete(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15), source='MANUAL',
                   review_status='APPROVED', approved_through_day=30)
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_COMPLETE)

    def test_manual_approved_through_day_twenty_is_partial(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15), source='MANUAL',
                   review_status='APPROVED', approved_through_day=20)
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_PARTIAL)

    def test_manual_needs_review_is_not_exempt(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15), source='MANUAL',
                   review_status='NEEDS_REVIEW', approved_through_day=30)
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_PARTIAL)

    def test_no_cards_stays_none(self):
        emp = self._employee()
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_NONE)

    # ---- Rule 3: user marked the received cards as the full submission ----
    def test_marked_complete_single_card_is_complete(self):
        emp = self._employee()
        # Mid-month card would otherwise be PARTIAL (Rule 1 doesn't reach it).
        self._card(emp, _utc(2026, 6, 15), completes_month=True)
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_COMPLETE)

    def test_marked_complete_does_not_require_approval(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15), review_status='NEEDS_REVIEW', completes_month=True)
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_COMPLETE)

    def test_marked_complete_on_one_of_several_cards_settles_the_month(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15))
        self._card(emp, _utc(2026, 6, 16), completes_month=True)
        db.session.commit()
        self.assertEqual(self._status(emp), mcs.STATUS_COMPLETE)

    def test_marked_complete_does_not_leak_into_another_month(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 5, 15), completes_month=True,
                   processing_month=date(2026, 5, 1))
        self._card(emp, _utc(2026, 6, 15))
        db.session.commit()
        # June has its own un-flagged card -> still a gap.
        self.assertEqual(self._status(emp), mcs.STATUS_PARTIAL)

    def test_marked_complete_row_is_flagged_and_counted_separately(self):
        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15), completes_month=True)
        gap = self._employee()
        self._card(gap, _utc(2026, 6, 15))
        db.session.commit()

        rows = mcs.compute_missing(self.business.id, MONTH, today=AFTER_MONTH_END)
        by_id = {r['employee_id']: r for r in rows}
        self.assertTrue(by_id[str(emp.id)]['is_exempt'])
        self.assertFalse(by_id[str(gap.id)]['is_exempt'])

        counts = mcs._bucket_counts(rows)
        self.assertEqual(counts['exempt'], 1)
        self.assertEqual(counts['missing'], 1)

        grp = mcs.group_by_site(rows)[0]
        self.assertEqual([e['employee_id'] for e in grp['employees']], [str(gap.id)])
        self.assertEqual([e['employee_id'] for e in grp['exempt_employees']], [str(emp.id)])
        self.assertEqual(grp['exempt_count'], 1)

    def test_manager_group_survives_when_every_gap_is_ignored(self):
        # Otherwise the group vanishes and the ignore can never be undone.
        manager = User(
            business_id=self.business.id,
            full_name='Mgr',
            email=f'mgr-{uuid.uuid4().hex[:8]}@example.com',
            password_hash='x',
            role='FIELD_MANAGER',
        )
        db.session.add(manager)
        db.session.flush()
        self.site.field_manager_id = manager.id

        emp = self._employee()
        self._card(emp, _utc(2026, 6, 15), completes_month=True)
        db.session.commit()

        rows = mcs.compute_missing(self.business.id, MONTH, today=AFTER_MONTH_END)
        groups = mcs.group_by_field_manager(rows)
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]['missing_count'], 0)
        self.assertEqual(groups[0]['exempt_count'], 1)

    # ---- report / bucket wiring reflects the exemption ----
    def test_exempted_employee_dropped_from_report_and_counts(self):
        exempt = self._employee()   # single late card -> COMPLETE
        self._card(exempt, _utc(2026, 6, 30))
        gap = self._employee()      # single mid-month card -> PARTIAL
        self._card(gap, _utc(2026, 6, 15))
        db.session.commit()

        rows = mcs.compute_missing(self.business.id, MONTH, today=AFTER_MONTH_END)
        by_mgr = mcs.group_by_site(rows)
        self.assertEqual(len(by_mgr), 1)
        grp = by_mgr[0]
        # Two employees on the site; one exempt (complete), one still a gap.
        self.assertEqual(grp['total_employees'], 2)
        self.assertEqual(grp['complete_count'], 1)
        self.assertEqual(grp['missing_count'], 1)
        listed = {e['employee_id'] for e in grp['employees']}
        self.assertEqual(listed, {str(gap.id)})

        # The Excel report lists only the remaining gap, not the exempt employee.
        from openpyxl import load_workbook
        gap_rows = [r for r in rows if r['status'] != mcs.STATUS_COMPLETE]
        output = mcs.generate_missing_cards_xlsx('T', gap_rows, MONTH)
        ws = load_workbook(output).active
        body = [[c.value for c in row] for row in ws.iter_rows(min_row=3)]
        self.assertEqual(len(body), 1)


if __name__ == '__main__':
    unittest.main()
