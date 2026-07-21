import unittest
import uuid
from copy import copy
from datetime import date
from types import SimpleNamespace

from openpyxl import Workbook

from backend.app.api.sites import LONG_DAY_FILL, _populate_template_core_sheet


def _is_grey(cell):
    return cell.fill is not None and cell.fill.fgColor.rgb == LONG_DAY_FILL.fgColor.rgb


class SummaryExportLongDayTests(unittest.TestCase):
    def _blank_sheet(self):
        workbook = Workbook()
        ws = workbook.active
        base_style = copy(ws['A1']._style)
        return workbook, ws, base_style

    def test_core_sheet_marks_long_days_grey(self):
        workbook, ws, base_style = self._blank_sheet()

        employee = SimpleNamespace(id=uuid.uuid4(), passport_id='P-300', full_name='Alice Doe')
        # Feb 2026 days 1-3 are Sun/Mon/Tue (not Saturday fallbacks).
        matrix = {
            str(employee.id): {1: 11.5, 2: 12.0, 3: 12.5},
        }
        status_matrix = {
            str(employee.id): {4: 'VACATION'},
        }

        _populate_template_core_sheet(
            ws=ws,
            employees=[employee],
            matrix=matrix,
            month_date=date(2026, 2, 1),
            style_header=base_style,
            style_body=base_style,
            style_total=base_style,
            status_matrix=status_matrix,
        )

        # Layout is transposed: employee at column 2, day N at row N+2.
        self.assertFalse(_is_grey(ws.cell(row=3, column=2)))  # day 1 = 11.5, below threshold
        self.assertTrue(_is_grey(ws.cell(row=4, column=2)))   # day 2 = 12.0, at threshold
        self.assertTrue(_is_grey(ws.cell(row=5, column=2)))   # day 3 = 12.5, above threshold

        vacation_cell = ws.cell(row=6, column=2)              # day 4 = VACATION label
        self.assertEqual(vacation_cell.value, 'חופשה')
        self.assertFalse(_is_grey(vacation_cell))


if __name__ == '__main__':
    unittest.main()
