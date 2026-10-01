"""
Absences reporting service.

Lists, per month, the employees who missed required workdays so the
organization can chase doctor approvals for sick leave early.

A *workday* is any Sunday–Friday day of the month (Saturday never counts) that
is not in the month's ``ignored_days`` setting and is not in the future —
during the current month only days up to and including today (Israel time)
are counted. Each employee may additionally have their own ignored days for
the month (e.g. before they started or after they left).

On a workday an employee is counted as missed when their merged day entry is:
    SICK   -> ``day_status == 'SICK'``
    EMPTY  -> no entry at all, or an entry with no hours, times or status
Anything else (hours worked, VACATION, HOLIDAY, INTERNATIONAL_VISA) is fine.

Only employees with at least one work card for the month are considered —
employees with no card at all are the Missing Cards report's concern.
"""
from __future__ import annotations

import calendar
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..extensions import db
from ..models.absences import AbsenceEmployeeDaySettings, AbsenceEmployeeExclusion, AbsenceMonthSettings
from ..models.sites import Employee, Site
from ..models.work_cards import WorkCard, WorkCardDayEntry, day_entry_has_data
from .missing_cards_service import _LOCAL_TZ

SATURDAY = 5  # date.weekday()

# Cards in these states don't hold usable day data for the month.
_IGNORED_CARD_STATUSES = ('REJECTED', 'SPLITTING')

CARD_APPROVED = 'APPROVED'
CARD_PENDING = 'PENDING'


def local_today() -> date:
    return datetime.now(_LOCAL_TZ).date()


def last_day_of_month(month: date) -> int:
    return calendar.monthrange(month.year, month.month)[1]


def normalize_ignored_days(month: date, days: Iterable[Any]) -> List[int]:
    """Validate and de-duplicate day numbers for ``month``. Raises ValueError."""
    last = last_day_of_month(month)
    result = set()
    for raw in days:
        if isinstance(raw, int) and not isinstance(raw, bool):
            day = raw
        elif isinstance(raw, str) and raw.strip().isdigit():
            day = int(raw)
        else:
            raise ValueError(f"Invalid day: {raw!r}")
        if not 1 <= day <= last:
            raise ValueError(f"Day {day} is outside {month.strftime('%Y-%m')}")
        result.add(day)
    return sorted(result)


def cutoff_day(month: date, today: date) -> int:
    """Last day-of-month that may be counted (0 for a future month)."""
    month_start = date(month.year, month.month, 1)
    if today < month_start:
        return 0
    if (today.year, today.month) == (month.year, month.month):
        return today.day
    return last_day_of_month(month)


def compute_workdays(month: date, ignored_days: Iterable[int], today: date) -> List[int]:
    """Day numbers that count as required workdays for ``month``."""
    ignored = set(ignored_days or [])
    last = cutoff_day(month, today)
    return [
        d for d in range(1, last + 1)
        if d not in ignored and date(month.year, month.month, d).weekday() != SATURDAY
    ]


def get_settings(business_id: UUID, month: date) -> Optional[AbsenceMonthSettings]:
    return (
        db.session.query(AbsenceMonthSettings)
        .filter(
            AbsenceMonthSettings.business_id == business_id,
            AbsenceMonthSettings.processing_month == month,
        )
        .first()
    )


def get_ignored_days(business_id: UUID, month: date) -> List[int]:
    settings = get_settings(business_id, month)
    return sorted(settings.ignored_days or []) if settings else []


def get_employee_day_settings(
    business_id: UUID, month: date, employee_id: UUID,
) -> Optional[AbsenceEmployeeDaySettings]:
    return (
        db.session.query(AbsenceEmployeeDaySettings)
        .filter(
            AbsenceEmployeeDaySettings.business_id == business_id,
            AbsenceEmployeeDaySettings.processing_month == month,
            AbsenceEmployeeDaySettings.employee_id == employee_id,
        )
        .first()
    )


def get_employee_ignored_days(business_id: UUID, month: date) -> Dict[UUID, List[int]]:
    """employee_id -> that employee's own ignored days for ``month``."""
    return {
        s.employee_id: sorted(s.ignored_days)
        for s in db.session.query(AbsenceEmployeeDaySettings).filter(
            AbsenceEmployeeDaySettings.business_id == business_id,
            AbsenceEmployeeDaySettings.processing_month == month,
        )
        if s.ignored_days
    }


def _entry_outranks(candidate: Dict[str, Any], existing: Dict[str, Any]) -> bool:
    """Per-day preference when merging an employee's cards: an APPROVED card's
    value beats a non-approved one; within a tier a day with data beats a blank
    one, then the newest card wins. Mirrors the review screen's consolidation."""
    if candidate['is_approved'] != existing['is_approved']:
        return candidate['is_approved']
    if candidate['has_data'] != existing['has_data']:
        return candidate['has_data']
    ca, ea = candidate['created_at'], existing['created_at']
    if ca is None or ea is None:
        return False
    return ca > ea


def classify_days(entries_by_day: Dict[int, Any], workdays: Iterable[int]) -> Dict[str, List[int]]:
    """Split ``workdays`` into sick / empty days given the merged entries."""
    sick: List[int] = []
    empty: List[int] = []
    for day in workdays:
        entry = entries_by_day.get(day)
        if entry is not None and entry.day_status == 'SICK':
            sick.append(day)
        elif not day_entry_has_data(entry):
            empty.append(day)
    return {'sick_days': sick, 'empty_days': empty}


def compute_absences(
    business_id: UUID,
    month: date,
    site_ids: Optional[List[UUID]] = None,
    today: Optional[date] = None,
) -> Dict[str, Any]:
    """Absences report for a business+month.

    ``site_ids`` narrows by the employee's home site (field-manager scoping);
    None means the whole business.
    """
    today = today or local_today()
    ignored_days = get_ignored_days(business_id, month)
    workdays = compute_workdays(month, ignored_days, today)
    employee_ignored = get_employee_ignored_days(business_id, month)

    excluded = {
        ex.employee_id: ex
        for ex in db.session.query(AbsenceEmployeeExclusion).filter(
            AbsenceEmployeeExclusion.business_id == business_id,
            AbsenceEmployeeExclusion.processing_month == month,
        )
    }

    card_query = (
        db.session.query(WorkCard)
        .join(Employee, Employee.id == WorkCard.employee_id)
        .filter(
            WorkCard.business_id == business_id,
            WorkCard.processing_month == month,
            WorkCard.employee_id.isnot(None),
            WorkCard.review_status.notin_(_IGNORED_CARD_STATUSES),
            Employee.business_id == business_id,
            Employee.is_active.is_(True),
        )
    )
    if site_ids is not None:
        card_query = card_query.filter(Employee.site_id.in_(site_ids))
    cards = card_query.all()

    entries_by_card: Dict[UUID, List[WorkCardDayEntry]] = {}
    card_ids = [c.id for c in cards]
    if card_ids:
        for entry in db.session.query(WorkCardDayEntry).filter(
            WorkCardDayEntry.work_card_id.in_(card_ids)
        ):
            entries_by_card.setdefault(entry.work_card_id, []).append(entry)

    cards_by_employee: Dict[UUID, List[WorkCard]] = {}
    for card in cards:
        cards_by_employee.setdefault(card.employee_id, []).append(card)

    employees: Dict[UUID, Any] = {}
    if cards_by_employee:
        emp_rows = (
            db.session.query(
                Employee.id, Employee.full_name, Employee.passport_id,
                Employee.external_employee_id, Site.id.label('site_id'), Site.site_name,
            )
            .outerjoin(Site, Site.id == Employee.site_id)
            .filter(Employee.id.in_(list(cards_by_employee.keys())))
        )
        employees = {r.id: r for r in emp_rows}

    rows: List[Dict[str, Any]] = []
    excluded_rows: List[Dict[str, Any]] = []
    # Every in-scope employee with personal ignored days, including those who
    # end up with no absences — so the UI can still show and undo them.
    employee_overrides: List[Dict[str, Any]] = []
    skipped_no_day_data = 0

    for employee_id, emp_cards in cards_by_employee.items():
        emp = employees.get(employee_id)
        if emp is None:
            continue

        own_ignored = employee_ignored.get(employee_id, [])
        if own_ignored:
            employee_overrides.append({
                'employee_id': str(employee_id),
                'full_name': emp.full_name,
                'site_name': emp.site_name,
                'ignored_days': own_ignored,
            })

        merged: Dict[int, Dict[str, Any]] = {}
        for card in emp_cards:
            is_approved = card.review_status == 'APPROVED'
            for entry in entries_by_card.get(card.id, []):
                candidate = {
                    'entry': entry,
                    'is_approved': is_approved,
                    'has_data': day_entry_has_data(entry),
                    'created_at': card.created_at,
                }
                existing = merged.get(entry.day_of_month)
                if existing is None or _entry_outranks(candidate, existing):
                    merged[entry.day_of_month] = candidate

        # A card with no day rows yet (extraction still running, or a single
        # monthly-total figure) would read as "every day empty" — skip it.
        if not merged:
            skipped_no_day_data += 1
            continue

        emp_workdays = [d for d in workdays if d not in own_ignored] if own_ignored else workdays
        classified = classify_days({d: c['entry'] for d, c in merged.items()}, emp_workdays)
        missed_total = len(classified['sick_days']) + len(classified['empty_days'])
        if missed_total == 0:
            continue

        ordered_cards = sorted(
            emp_cards,
            key=lambda c: (c.review_status != 'APPROVED', -(c.created_at.timestamp() if c.created_at else 0)),
        )
        row = {
            'employee_id': str(employee_id),
            'full_name': emp.full_name,
            'passport_id': emp.passport_id,
            'external_employee_id': emp.external_employee_id,
            'site_id': str(emp.site_id) if emp.site_id else None,
            'site_name': emp.site_name,
            'sick_days': classified['sick_days'],
            'empty_days': classified['empty_days'],
            'sick_count': len(classified['sick_days']),
            'empty_count': len(classified['empty_days']),
            'missed_total': missed_total,
            'ignored_days': own_ignored,
            # APPROVED only once every card is approved — an unreviewed card
            # means the month's data may still change.
            'card_status': CARD_APPROVED if all(c.review_status == 'APPROVED' for c in emp_cards) else CARD_PENDING,
            'card_ids': [str(c.id) for c in ordered_cards],
            # The card to open for review: approved first, then newest.
            'review_card_id': str(ordered_cards[0].id),
            'review_site_id': str(ordered_cards[0].site_id) if ordered_cards[0].site_id else None,
        }
        exclusion = excluded.get(employee_id)
        if exclusion is not None:
            row['exclusion_reason'] = exclusion.reason
            excluded_rows.append(row)
        else:
            rows.append(row)

    sort_key = lambda r: (-r['missed_total'], r['site_name'] or '', r['full_name'] or '')
    rows.sort(key=sort_key)
    excluded_rows.sort(key=sort_key)
    employee_overrides.sort(key=lambda o: (o['site_name'] or '', o['full_name'] or ''))

    return {
        'month': month.isoformat(),
        'settings': {'ignored_days': ignored_days},
        'cutoff_day': cutoff_day(month, today),
        'workdays': workdays,
        'workdays_counted': len(workdays),
        'summary': {
            'employees': len(rows),
            'sick_days': sum(r['sick_count'] for r in rows),
            'empty_days': sum(r['empty_count'] for r in rows),
            'excluded': len(excluded_rows),
            'skipped_no_day_data': skipped_no_day_data,
        },
        'rows': rows,
        'excluded_rows': excluded_rows,
        'employee_overrides': employee_overrides,
    }


_HEADERS = ['מספר עובד', 'שם עובד', 'דרכון', 'אתר', 'ימי מחלה', 'ימים', 'ימים ללא דיווח', 'ימים', 'סה"כ', 'סטטוס כרטיס']
_WIDTHS = [14, 22, 16, 20, 10, 28, 14, 28, 8, 14]
_CARD_STATUS_HE = {CARD_APPROVED: 'מאושר', CARD_PENDING: 'ממתין לאישור'}


def _days_label(month: date, days: List[int]) -> str:
    return ', '.join(f"{d:02d}/{month.month:02d}" for d in days)


def generate_absences_xlsx(result: Dict[str, Any], month: date, title: str) -> BytesIO:
    """Single-sheet RTL XLSX of the (non-excluded) absence rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = month.strftime('%Y-%m')
    ws.sheet_view.rightToLeft = True

    ws.append([f"היעדרויות — {title} — {month.strftime('%Y-%m')}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='right')

    ws.append(_HEADERS)
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    for col in range(1, len(_HEADERS) + 1):
        cell = ws.cell(row=2, column=col)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='right')

    for row in result['rows']:
        ws.append([
            row.get('external_employee_id') or '',
            row.get('full_name') or '',
            row.get('passport_id') or '',
            row.get('site_name') or '',
            row['sick_count'],
            _days_label(month, row['sick_days']),
            row['empty_count'],
            _days_label(month, row['empty_days']),
            row['missed_total'],
            _CARD_STATUS_HE.get(row['card_status'], row['card_status']),
        ])

    for i, w in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
